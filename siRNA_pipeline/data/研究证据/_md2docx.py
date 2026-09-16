# -*- coding: utf-8 -*-
"""Markdown -> DOCX converter for the mismatch-research reports.

python-docx has no Markdown reader, so this parses the subset these reports use:
ATX headings, fenced code blocks, GitHub pipe tables, bullet / ordered lists,
blockquotes, horizontal rules, and the inline spans **bold**, *italic*,
`code`, and [text](url).  Every run carries an eastAsia font so Chinese
renders with a proper CJK face instead of a fallback box.

Usage:  python _md2docx.py <file.md> [<file.md> ...]
Writes <file>.docx next to each source and prints a one-line report per file.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

CJK = "微软雅黑"
MONO = "Consolas"

HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
HR_RE = re.compile(r"^\s{0,3}(?:-{3,}|\*{3,}|_{3,})\s*$")
BULLET_RE = re.compile(r"^(\s*)[-*+]\s+(.*)$")
ORDERED_RE = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
FENCE_RE = re.compile(r"^\s*(?:```|~~~)\s*[\w+-]*\s*$")
TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(?:\|\s*:?-{2,}:?\s*)*\|?\s*$")
INLINE_RE = re.compile(
    r"(\*\*[^*]+?\*\*|`[^`]+`|(?<![\w*])\*[^*\n]+?\*(?![\w*])|\[[^\]]+\]\([^)\s]+\))"
)


def style_run(run, *, bold=None, italic=False, mono=False, size=None, color=None):
    """Apply a CJK-safe font plus optional weight/colour to one run."""
    name = MONO if mono else CJK
    run.font.name = name
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), CJK)


def add_inline(paragraph, text, size=None):
    """Emit one paragraph's worth of inline Markdown spans."""
    text = text.strip()
    if not text:
        return
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            style_run(paragraph.add_run(text[pos:match.start()]), size=size)
        token = match.group(0)
        if token.startswith("**"):
            style_run(paragraph.add_run(token[2:-2]), bold=True, size=size)
        elif token.startswith("`"):
            style_run(paragraph.add_run(token[1:-1]), mono=True, size=9)
        elif token.startswith("["):
            label, _, url = token[1:-1].partition("](")
            style_run(paragraph.add_run(label), size=size)
            if url and url != label:
                style_run(paragraph.add_run(f" ({url})"), size=7.5,
                          color=RGBColor(0x60, 0x60, 0x60))
        else:
            style_run(paragraph.add_run(token[1:-1]), italic=True, size=size)
        pos = match.end()
    if pos < len(text):
        style_run(paragraph.add_run(text[pos:]), size=size)


def split_row(line):
    """Split a pipe-table row into trimmed cells."""
    body = line.strip()
    if body.startswith("|"):
        body = body[1:]
    if body.endswith("|"):
        body = body[:-1]
    return [cell.strip() for cell in body.split("|")]


def convert(src: Path) -> Path:
    lines = src.read_text(encoding="utf-8").splitlines()
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = CJK
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)

    index = 0
    total = len(lines)
    while index < total:
        line = lines[index]

        # fenced code block -------------------------------------------------
        if FENCE_RE.match(line):
            index += 1
            block = []
            while index < total and not FENCE_RE.match(lines[index]):
                block.append(lines[index])
                index += 1
            index += 1
            for code_line in block:
                para = doc.add_paragraph()
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.left_indent = Inches(0.25)
                style_run(para.add_run(code_line if code_line else " "),
                          mono=True, size=8.5)
            continue

        # pipe table --------------------------------------------------------
        if (line.strip().startswith("|") and index + 1 < total
                and TABLE_SEP_RE.match(lines[index + 1])):
            rows = []
            while index < total and lines[index].strip().startswith("|"):
                rows.append(split_row(lines[index]))
                index += 1
            rows.pop(1)  # drop the --- separator row
            if rows:
                width = max(len(r) for r in rows)
                table = doc.add_table(rows=0, cols=width)
                table.style = "Table Grid"
                for row_cells in rows:
                    cells = table.add_row().cells
                    for col in range(width):
                        text = row_cells[col] if col < len(row_cells) else ""
                        para = cells[col].paragraphs[0]
                        add_inline(para, text, size=9)
            doc.add_paragraph()
            continue

        # heading -----------------------------------------------------------
        heading = HEADING_RE.match(line)
        if heading:
            level = min(len(heading.group(1)), 6)
            para = doc.add_heading(level=level)
            add_inline(para, heading.group(2))
            index += 1
            continue

        # horizontal rule ---------------------------------------------------
        if HR_RE.match(line):
            index += 1
            continue

        # blockquote --------------------------------------------------------
        if line.lstrip().startswith(">"):
            block = []
            while index < total and lines[index].lstrip().startswith(">"):
                block.append(lines[index].lstrip()[1:].strip())
                index += 1
            para = doc.add_paragraph(style="Intense Quote")
            add_inline(para, " ".join(part for part in block if part), size=10)
            continue

        # lists -------------------------------------------------------------
        bullet = BULLET_RE.match(line)
        ordered = ORDERED_RE.match(line)
        if bullet or ordered:
            match = bullet or ordered
            indent = len(match.group(1)) // 2
            text = match.group(3) if ordered else match.group(2)
            style = "List Number" if ordered else "List Bullet"
            para = doc.add_paragraph(style=style)
            para.paragraph_format.left_indent = Inches(0.25 * (indent + 1))
            add_inline(para, text)
            index += 1
            continue

        # blank / paragraph --------------------------------------------------
        if not line.strip():
            index += 1
            continue
        para = doc.add_paragraph()
        add_inline(para, line)
        index += 1

    out = src.with_suffix(".docx")
    doc.save(str(out))
    return out


def main(argv):
    if not argv:
        print("usage: python _md2docx.py <file.md> [...]")
        return 2
    for raw in argv:
        src = Path(raw)
        if not src.is_file():
            print(f"SKIP (missing): {src}")
            continue
        out = convert(src)
        paras = len(Document(str(out)).paragraphs)
        print(f"OK  {src.name} -> {out.name}  "
              f"({out.stat().st_size // 1024} KB, {paras} paragraphs)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
